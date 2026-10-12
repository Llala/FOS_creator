"""Parser for curriculum XLSX files exported by 1C University/Plans.

The parser normalizes data needed by a FOS generation pipeline:
program metadata, disciplines, semester workload, assessment forms and
competencies. It does not use an LLM, so values from the curriculum remain
deterministic and auditable.

Usage:
    python fos_curriculum_parser.py curriculum.xlsx -o curriculum.json
    python fos_curriculum_parser.py curriculum.xlsx --discipline-code Б1.В.07
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


ASSESSMENT_COLUMNS = {
    4: "exam",
    5: "pass",
    6: "graded_pass",
    7: "course_project",
    8: "course_work",
}

# In the "План" sheet each semester occupies seven columns.
SEMESTER_COLUMNS = {
    semester: {
        "credits": 17 + (semester - 1) * 7,
        "total_hours": 18 + (semester - 1) * 7,
        "lectures": 19 + (semester - 1) * 7,
        "laboratories": 20 + (semester - 1) * 7,
        "practicals": 21 + (semester - 1) * 7,
        "independent_work": 22 + (semester - 1) * 7,
        "control_hours": 23 + (semester - 1) * 7,
    }
    for semester in range(1, 9)
}

COMPONENT_CODE_RE = re.compile(
    r"^(?:Б[1-3]|ФТД)(?:[.()]|$)", re.IGNORECASE
)
COMPETENCY_CODE_RE = re.compile(r"^(?:УК|ОПК|ПК)-\d+(?:\.\d+)?$", re.IGNORECASE)


class CurriculumFormatError(ValueError):
    """Raised when a workbook does not have the expected curriculum sheets."""


@dataclass(slots=True)
class Competency:
    code: str
    description: str | None = None
    competency_type: str | None = None


@dataclass(slots=True)
class Assessment:
    kind: str
    semester: int


@dataclass(slots=True)
class SemesterWorkload:
    semester: int
    credits: float | None = None
    total_hours: float | None = None
    lectures: float | None = None
    laboratories: float | None = None
    practicals: float | None = None
    independent_work: float | None = None
    control_hours: float | None = None


@dataclass(slots=True)
class CurriculumComponent:
    source_row: int
    code: str
    name: str
    component_type: str
    credits_total: float | None
    hours_total: float | None
    classroom_hours: float | None
    independent_work_total: float | None
    control_hours_total: float | None
    department_code: str | None
    department_name: str | None
    assessments: list[Assessment] = field(default_factory=list)
    competencies: list[Competency] = field(default_factory=list)
    semesters: list[SemesterWorkload] = field(default_factory=list)


@dataclass(slots=True)
class ProgramMetadata:
    direction_code: str | None = None
    direction_name: str | None = None
    profile: str | None = None
    qualification: str | None = None
    study_form: str | None = None
    duration: str | None = None
    start_years: list[int] = field(default_factory=list)
    educational_standard: str | None = None
    graduating_department: str | None = None


@dataclass(slots=True)
class Curriculum:
    source_file: str
    program: ProgramMetadata
    components: list[CurriculumComponent]

    def find_components(self, code: str) -> list[CurriculumComponent]:
        normalized = normalize_text(code)
        return [
            component
            for component in self.components
            if normalize_text(component.code) == normalized
        ]

    def find_component(self, code: str) -> CurriculumComponent:
        matches = self.find_components(code)
        if not matches:
            raise KeyError(f"Component {code!r} was not found")
        if len(matches) > 1:
            choices = "; ".join(
                f"row {item.source_row}: {item.name}" for item in matches
            )
            raise KeyError(
                f"Component code {code!r} is ambiguous. Matches: {choices}"
            )
        return matches[0]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def normalize_text(value: Any) -> str:
    """Collapse whitespace and convert an Excel value to clean text."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def number_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value).replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def int_text_or_none(value: Any) -> str | None:
    """Keep department identifiers stable instead of returning 13.0."""
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value) or None


def split_codes(value: Any) -> list[str]:
    """Parse semicolon/comma separated competency codes."""
    text = normalize_text(value)
    if not text:
        return []
    result: list[str] = []
    for token in re.split(r"[;,\n]+", text):
        code = normalize_text(token).upper()
        if code and COMPETENCY_CODE_RE.match(code) and code not in result:
            result.append(code)
    return result


def parse_semester_numbers(value: Any) -> list[int]:
    """Read assessment semesters such as 4, '1; 7' or legacy '17'."""
    if value in (None, ""):
        return []
    if isinstance(value, (int, float)):
        text = str(int(value))
    else:
        text = normalize_text(value)

    explicit = [int(x) for x in re.findall(r"(?<!\d)[1-8](?!\d)", text)]
    if explicit:
        return list(dict.fromkeys(explicit))

    # Some curriculum exports concatenate one-digit semesters: 17 = 1 and 7.
    compact = re.sub(r"\D", "", text)
    if compact and all(ch in "12345678" for ch in compact):
        return list(dict.fromkeys(int(ch) for ch in compact))
    return []


def component_type(code: str) -> str:
    upper = code.upper()
    if upper.startswith("Б1"):
        return "discipline"
    if upper.startswith("Б2"):
        return "practice"
    if upper.startswith("Б3"):
        return "state_final_assessment"
    if upper.startswith("ФТД"):
        return "elective"
    return "other"


def adjacent_value(ws: Worksheet, row: int, column: int) -> str | None:
    """Find the next non-empty cell on the same row."""
    for next_column in range(column + 1, ws.max_column + 1):
        value = normalize_text(ws.cell(row, next_column).value)
        if value:
            return value
    return None


def parse_labeled_inline(text: str, label: str) -> str | None:
    match = re.search(re.escape(label) + r"\s*:?[ ]*(.+)$", text, re.IGNORECASE)
    return normalize_text(match.group(1)) if match else None


def parse_program_metadata(ws: Worksheet) -> ProgramMetadata:
    metadata = ProgramMetadata()

    for row in ws.iter_rows():
        for cell in row:
            text = normalize_text(cell.value)
            if not text:
                continue

            if text.startswith("Направление подготовки"):
                match = re.search(
                    r"Направление подготовки\s+(\d{2}\.\d{2}\.\d{2})\s+(.+)", text
                )
                if match:
                    metadata.direction_code = match.group(1)
                    metadata.direction_name = normalize_text(match.group(2))
            elif text.rstrip(":") == "Профиль":
                metadata.profile = adjacent_value(ws, cell.row, cell.column)
            elif text.startswith("Кафедра:"):
                metadata.graduating_department = adjacent_value(ws, cell.row, cell.column)
            elif text.startswith("Квалификация:"):
                metadata.qualification = parse_labeled_inline(text, "Квалификация")
            elif text.startswith("Форма обучения:"):
                metadata.study_form = parse_labeled_inline(text, "Форма обучения")
            elif text.startswith("Срок получения образования:"):
                metadata.duration = parse_labeled_inline(text, "Срок получения образования")
            elif text == "Год начала подготовки (по учебному плану)":
                years = adjacent_value(ws, cell.row, cell.column) or ""
                metadata.start_years = [int(x) for x in re.findall(r"20\d{2}", years)]
            elif text == "Образовательный стандарт (ФГОС)":
                metadata.educational_standard = adjacent_value(ws, cell.row, cell.column)

    return metadata


def parse_competency_catalog(ws: Worksheet) -> dict[str, Competency]:
    catalog: dict[str, Competency] = {}
    for row in ws.iter_rows(values_only=True):
        code = normalize_text(row[1] if len(row) > 1 else None).upper()
        description = normalize_text(row[3] if len(row) > 3 else None)
        kind = normalize_text(row[4] if len(row) > 4 else None)
        if COMPETENCY_CODE_RE.match(code) and description:
            catalog[code] = Competency(
                code=code,
                description=description,
                competency_type=kind or code.split("-")[0],
            )
    return catalog


def parse_semesters(row: tuple[Any, ...]) -> list[SemesterWorkload]:
    workloads: list[SemesterWorkload] = []
    for semester, columns in SEMESTER_COLUMNS.items():
        values = {
            field_name: number_or_none(row[column - 1])
            for field_name, column in columns.items()
        }
        if any(value is not None for value in values.values()):
            workloads.append(SemesterWorkload(semester=semester, **values))
    return workloads


def infer_assessment_semesters(
    assessment_cell: Any, active_semesters: Iterable[int]
) -> list[int]:
    parsed = parse_semester_numbers(assessment_cell)
    if parsed:
        return parsed
    # Defensive fallback for exports where the assessment marker is '+' or 'x'.
    if normalize_text(assessment_cell):
        return list(active_semesters)
    return []


def parse_components(
    ws: Worksheet,
    competency_catalog: dict[str, Competency],
) -> list[CurriculumComponent]:
    components: list[CurriculumComponent] = []

    for source_row, row in enumerate(
        ws.iter_rows(min_row=4, values_only=True), start=4
    ):
        code = normalize_text(row[1] if len(row) > 1 else None)
        name = normalize_text(row[2] if len(row) > 2 else None)
        included = normalize_text(row[0] if row else None)
        if not code or not name or not COMPONENT_CODE_RE.match(code):
            continue
        if included and included not in {"+", "1", "да", "Да", "ДА"}:
            continue

        semesters = parse_semesters(row)
        active_semesters = [item.semester for item in semesters]

        assessments: list[Assessment] = []
        for column, kind in ASSESSMENT_COLUMNS.items():
            for semester in infer_assessment_semesters(row[column - 1], active_semesters):
                assessments.append(Assessment(kind=kind, semester=semester))

        competency_codes = split_codes(row[74] if len(row) >= 75 else None)
        competencies = [
            competency_catalog.get(code, Competency(code=code))
            for code in competency_codes
        ]

        components.append(
            CurriculumComponent(
                source_row=source_row,
                code=code,
                name=name,
                component_type=component_type(code),
                credits_total=number_or_none(row[9]),
                hours_total=number_or_none(row[12]),
                classroom_hours=number_or_none(row[13]),
                independent_work_total=number_or_none(row[14]),
                control_hours_total=number_or_none(row[15]),
                department_code=int_text_or_none(row[72] if len(row) >= 73 else None),
                department_name=normalize_text(row[73] if len(row) >= 74 else None) or None,
                assessments=assessments,
                competencies=competencies,
                semesters=semesters,
            )
        )

    return components


def parse_curriculum(path: str | Path) -> Curriculum:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)

    workbook = load_workbook(source, read_only=True, data_only=True)
    required_sheets = {"Титул", "План", "Компетенции"}
    missing = required_sheets - set(workbook.sheetnames)
    if missing:
        raise CurriculumFormatError(
            "Missing required sheets: " + ", ".join(sorted(missing))
        )

    try:
        catalog = parse_competency_catalog(workbook["Компетенции"])
        return Curriculum(
            source_file=source.name,
            program=parse_program_metadata(workbook["Титул"]),
            components=parse_components(workbook["План"], catalog),
        )
    finally:
        workbook.close()


def write_json(data: Any, path: str | Path) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def build_cli() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Parse a curriculum XLSX into normalized JSON for a FOS system."
    )
    parser.add_argument("xlsx", type=Path, help="Path to the curriculum XLSX")
    parser.add_argument("-o", "--output", type=Path, help="Destination JSON file")
    parser.add_argument(
        "--discipline-code",
        help="Return only one component, for example Б1.В.07",
    )
    return parser


def main() -> int:
    args = build_cli().parse_args()
    curriculum = parse_curriculum(args.xlsx)

    if args.discipline_code:
        try:
            result: Any = asdict(curriculum.find_component(args.discipline_code))
        except KeyError as error:
            raise SystemExit(str(error)) from error
    else:
        result = curriculum.to_dict()

    if args.output:
        write_json(result, args.output)
        print(f"Saved: {args.output}")
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
