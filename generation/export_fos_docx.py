"""Export a validated generate_fos.py result to an editable draft DOCX.

Usage: python export_fos_docx.py physical_culture_fos.json -o draft.docx --page-size a4
No model requests, embeddings or API key required.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from generate_fos import validate_fos


ASSESSMENTS = {"exam": "Экзамен", "pass": "Зачёт", "graded_pass": "Зачёт с оценкой",
               "course_work": "Курсовая работа", "course_project": "Курсовой проект"}
FORMATS = {"question": "Вопрос", "test": "Тест", "case": "Кейс",
           "practical": "Практическое задание", "project": "Проект"}


def label(doc, title: str, text: Any):
    paragraph = doc.add_paragraph()
    paragraph.add_run(title + ": ").bold = True
    paragraph.add_run(str(text))
    return paragraph


def body(doc, text: str):
    # Preserve user/model line breaks; never interpret the text as executable markup.
    return doc.add_paragraph(text)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float]):
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    properties = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)
    repeat_header = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for cell, text in zip(result.rows[0].cells, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        for cell, text in zip(result.add_row().cells, values):
            cell.text = str(text)
    for row in result.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    run.font.size = Pt(11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return result


def validate_export(result: dict[str, Any]):
    if not isinstance(result, dict):
        raise ValueError("Ожидается полный JSON-результат generate_fos.py")
    if result.get("status") != "requires_teacher_review":
        raise ValueError("Экспорт разрешён для status=requires_teacher_review; сначала исправьте генерацию")
    if result.get("validation", {}).get("valid") is not True:
        raise ValueError("В исходном результате validation.valid должен быть true")
    context = {"curriculum_facts": result.get("curriculum_facts"),
               "sources": result.get("sources"), "missing_inputs": []}
    fos, report = validate_fos(result.get("fos"), context, min_tasks=1)
    if not report["valid"]:
        details = "; ".join(e["path"] + ": " + e["message"] for e in report["errors"])
        raise ValueError("Повторная проверка ФОС не пройдена: " + details)
    return fos


def export_fos(result: dict[str, Any], output: str | Path, page_size: str = "letter") -> Path:
    fos = validate_export(result)
    path = Path(output)
    if path.suffix.lower() != ".docx":
        raise ValueError("Выходной файл должен иметь расширение .docx")
    if page_size not in {"a4", "letter"}:
        raise ValueError("page_size: a4 или letter")
    facts = result["curriculum_facts"]
    doc = Document()
    doc.core_properties.title = f"Черновик ФОС {fos.discipline_name}"
    doc.core_properties.subject = "Для проверки преподавателем"
    doc.core_properties.author = ""
    section = doc.sections[0]
    if page_size == "a4":
        section.page_width, section.page_height = Cm(21), Cm(29.7)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.add_paragraph("Черновик фонда оценочных средств", style="Title")
    doc.add_paragraph(fos.discipline_name, style="Subtitle")
    body(doc, "Документ предназначен для проверки преподавателем. Задания, ответы и "
         "критерии требуют содержательной проверки и согласования перед использованием.")

    doc.add_heading("Сведения о дисциплине", level=1)
    label(doc, "Код", fos.discipline_code)
    if facts.get("hours_total") is not None:
        label(doc, "Общая трудоёмкость в часах", facts["hours_total"])
    if facts.get("credits_total") is not None:
        label(doc, "Зачётные единицы", facts["credits_total"])
    for assessment in facts.get("assessments", []):
        label(doc, "Аттестация", f"{ASSESSMENTS.get(assessment['kind'], assessment['kind'])}, "
              f"семестр {assessment['semester']}")
    doc.add_heading("Компетенции", level=1)
    for competency in facts["competencies"]:
        label(doc, competency["code"], competency.get("description") or "Описание не указано")

    warnings = list(dict.fromkeys(
        result["validation"].get("warnings", []) + fos.missing_data))
    if warnings:
        doc.add_heading("Замечания для проверки", level=1)
        for warning in warnings:
            body(doc, str(warning))

    doc.add_heading("Оценочные задания", level=1)
    for number, task in enumerate(fos.tasks, 1):
        doc.add_heading(f"Задание {number}", level=2)
        label(doc, "Идентификатор", task.id)
        label(doc, "Тема", task.topic)
        label(doc, "Компетенции", ", ".join(task.competency_codes))
        label(doc, "Проверяемое умение", task.assessed_skill)
        label(doc, "Формат", FORMATS[task.format])
        control = "Текущий контроль" if task.stage == "current" else (
            f"{ASSESSMENTS.get(task.assessment_kind, task.assessment_kind)}, семестр {task.semester}")
        label(doc, "Контроль", control)
        label(doc, "Максимальный балл", task.max_points)
        doc.add_heading("Формулировка задания", level=3)
        body(doc, task.prompt)
        doc.add_heading("Ожидаемый ответ", level=3)
        body(doc, task.expected_answer)
        doc.add_heading("Критерии оценивания", level=3)
        width = section.page_width.cm - 4
        table(doc, ["Критерий", "Максимальный балл"],
              [[c.description, str(c.max_points)] for c in task.criteria], [width - 3.5, 3.5])
        for i, criterion in enumerate(task.criteria, 1):
            label(doc, f"Критерий {i}", criterion.description)
            label(doc, "Ноль баллов", criterion.zero_points)
            label(doc, f"Полный балл {criterion.max_points}", criterion.full_points)
        doc.add_heading("Основания в РПД и других источниках", level=3)
        for evidence in task.evidence:
            body(doc, f"[{evidence.ref}] {evidence.quote}")

    doc.add_heading("Источники", level=1)
    for source in result["sources"]:
        text = source.get("document") or source.get("source_id") or "Название не указано"
        details = [source.get("location"), source.get("section")]
        label(doc, f"[{source['ref']}]", text)
        for detail in details:
            if detail:
                body(doc, detail)

    # Complete the ZIP before replacing the destination file.
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        doc.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Полный JSON-результат generate_fos.py")
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--page-size", choices=["letter", "a4"], default="letter")
    args = parser.parse_args()
    try:
        if args.input.resolve() == args.output.resolve():
            raise ValueError("Входной JSON и выходной DOCX должны быть разными файлами")
        result = json.loads(args.input.read_text(encoding="utf-8-sig"))
        path = export_fos(result, args.output, args.page_size)
        print(f"Черновик сохранён: {path}")
        return 0
    except (ValueError, OSError, KeyError, TypeError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
