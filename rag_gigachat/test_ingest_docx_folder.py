from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from fos_rag import RAGStore
from ingest_docx_folder import (
    build_jobs, discover_docx, infer_code, ingest_jobs, load_curriculum,
)


class FakeEmbedder:
    model = "fake-v1"

    def __init__(self, _: str = "fake-v1") -> None:
        pass

    def embed(self, texts):
        return [[float(len(text)), 1.0] for text in texts]


def make_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_heading("Темы")
    document.add_paragraph(text)
    document.save(path)


class BatchIngestTests(unittest.TestCase):
    def test_disciplines_are_inferred_from_plan_or_reported_ambiguous(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            plan = folder / "plan.json"
            plan.write_text('''{"components": [
                {"code":"Б1.В.07","name":"Маркетинг","competencies":[{"code":"ПК-7"}]},
                {"code":"Б1.В.08","name":"Маркетинг спортивных услуг","competencies":[{"code":"ПК-9"}]}
            ]}''', encoding="utf-8")
            catalog = load_curriculum(plan)
            self.assertEqual(infer_code("РПД Б1.В.07 Маркетинг.docx", catalog), "Б1.В.07")
            self.assertEqual(infer_code("ФОС Маркетинг 2025.docx", catalog), "Б1.В.07")
            with self.assertRaisesRegex(ValueError, "неоднозначное"):
                infer_code("Материалы Маркетинг спортивных услуг дополнение.docx", catalog)
            rpd = folder / "РПД Маркетинг.docx"
            make_docx(rpd, "Анализ рынка")
            jobs = build_jobs(folder, [rpd], {}, "rpd", None, (), catalog)
            self.assertEqual((jobs[0].discipline_code, jobs[0].competencies),
                             ("Б1.В.07", ("ПК-7",)))

    def test_nested_curriculum_codes_select_the_full_code(self):
        catalog = {
            "Б1.В.ДВ.01": {"name": "Дисциплины по выбору"},
            "Б1.В.ДВ.01.01": {"name": "Речевые коммуникации в управленческой деятельности"},
            "Б1.В.ДВ.01.02": {"name": "Иная дисциплина"},
        }
        filename = "Б1.В.ДВ.01.01 Речевые коммуникации в управленческой деятельности РП.docx"
        self.assertEqual(infer_code(filename, catalog), "Б1.В.ДВ.01.01")
        self.assertEqual(infer_code("Б1.В.ДВ.01 Дисциплины по выбору.docx", catalog),
                         "Б1.В.ДВ.01")
        with self.assertRaisesRegex(ValueError, "несколько кодов"):
            infer_code("Б1.В.ДВ.01.01 и Б1.В.ДВ.01.02.docx", catalog)

    def test_foreign_plan_code_and_mismatched_title(self):
        catalog = {
            "Б1.В.ДВ.01": {"name": "Дисциплины по выбору", "competencies": []},
            "Б1.В.ДВ.01.01": {
                "name": "Международный студенческий спорт", "competencies": ["ПК-6"],
            },
        }
        with self.assertRaisesRegex(ValueError, "соответствует"):
            infer_code("Б1.В.ДВ.01.01 Речевые коммуникации в управленческой деятельности РП.docx", catalog)
        foreign = "Б1.В.ДВ.01.02 Речевое общение в менеджменте РП.docx"
        with self.assertRaisesRegex(ValueError, "отсутствует в учебном плане"):
            infer_code(foreign, catalog)
        self.assertEqual(infer_code(foreign, catalog, allow_unknown_codes=True),
                         "Б1.В.ДВ.01.02")

    def test_recursive_manifest_and_idempotent_ingest(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "docs"
            rpd_dir = folder / "rpd"
            rpd_dir.mkdir(parents=True)
            make_docx(rpd_dir / "marketing.docx", "Рынок спортивных услуг")
            make_docx(folder / "template.DOCX", "Критерии оценивания")
            (folder / "~$temporary.docx").write_bytes(b"")
            files = discover_docx(folder)
            self.assertEqual(len(files), 2)
            manifest = {
                "rpd/marketing.docx": {
                    "document_type": "rpd",
                    "discipline_code": "Б1.В.07",
                    "competencies": ["ПК-7"],
                },
                "template.DOCX": {"document_type": "fos_template"},
            }
            jobs = build_jobs(folder, files, manifest, None, None, ())
            db = Path(temp) / "index.sqlite3"
            result = ingest_jobs(jobs, db, "fake-v1", FakeEmbedder)
            self.assertEqual((result["indexed"], result["failed"]), (2, 0))
            again = ingest_jobs(jobs, db, "fake-v1", FakeEmbedder)
            self.assertEqual((again["unchanged"], again["failed"]), (2, 0))
            with RAGStore(db, FakeEmbedder()) as store:
                hits = store.search("рынок", "Б1.В.07")
            self.assertEqual({hit.document_name for hit in hits},
                             {"marketing.docx", "template.DOCX"})

    def test_prevalidation_rejects_missing_code_before_indexing(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            make_docx(folder / "rpd.docx", "Тема дисциплины")
            with self.assertRaisesRegex(ValueError, "discipline_code"):
                build_jobs(folder, discover_docx(folder), {}, "rpd", None, ())

    def test_error_in_one_file_does_not_stop_other_files(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            make_docx(folder / "good.docx", "Методические требования")
            (folder / "bad.docx").write_bytes(b"not a Word document")
            jobs = build_jobs(folder, discover_docx(folder), {}, "methodology", None, ())
            result = ingest_jobs(jobs, folder / "index.sqlite3", "fake-v1", FakeEmbedder)
            self.assertEqual((result["indexed"], result["failed"]), (1, 1))


if __name__ == "__main__":
    unittest.main()
